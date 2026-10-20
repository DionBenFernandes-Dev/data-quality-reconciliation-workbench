from __future__ import annotations

import csv
import io
import json
import os
import re
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

from .engine import is_missing, json_safe, reconcile_dataframes

APP_DIR = Path(__file__).resolve().parent
BACKEND_DIR = APP_DIR.parent
DATA_DIR = Path(os.getenv("DQW_DATA_DIR", BACKEND_DIR / "data")).resolve()
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = DATA_DIR / "workbench.sqlite3"
SAMPLE_DIR = BACKEND_DIR / "sample_data"
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 250_000
ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}

DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(
    title="Data Quality & Reconciliation Workbench API",
    version="0.1.0",
    description="Local-first, explainable reconciliation for CSV and Excel files.",
)
allowed_origins = [x.strip() for x in os.getenv(
    "CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173,http://localhost:4173,http://127.0.0.1:4173"
).split(",") if x.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=False,
    allow_methods=["GET", "POST", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type"],
)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def connect_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    return conn


def init_db() -> None:
    with connect_db() as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS uploads (
            id TEXT PRIMARY KEY,
            filename TEXT NOT NULL,
            stored_path TEXT NOT NULL,
            extension TEXT NOT NULL,
            selected_sheet TEXT,
            sheet_names_json TEXT NOT NULL DEFAULT '[]',
            columns_json TEXT NOT NULL,
            row_count INTEGER NOT NULL,
            missing_cells INTEGER NOT NULL,
            size_bytes INTEGER NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS runs (
            id TEXT PRIMARY KEY,
            created_at TEXT NOT NULL,
            name TEXT NOT NULL,
            source_a_upload_id TEXT NOT NULL,
            source_b_upload_id TEXT NOT NULL,
            source_a_name TEXT NOT NULL,
            source_b_name TEXT NOT NULL,
            summary_json TEXT NOT NULL,
            config_json TEXT NOT NULL,
            warnings_json TEXT NOT NULL DEFAULT '[]',
            FOREIGN KEY(source_a_upload_id) REFERENCES uploads(id),
            FOREIGN KEY(source_b_upload_id) REFERENCES uploads(id)
        );
        CREATE TABLE IF NOT EXISTS run_exceptions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id TEXT NOT NULL,
            ordinal INTEGER NOT NULL,
            exception_type TEXT NOT NULL,
            item_json TEXT NOT NULL,
            FOREIGN KEY(run_id) REFERENCES runs(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS ix_run_exceptions_run_type ON run_exceptions(run_id, exception_type, ordinal);
        CREATE TABLE IF NOT EXISTS profiles (
            id TEXT PRIMARY KEY,
            name TEXT NOT NULL UNIQUE COLLATE NOCASE,
            config_json TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at TEXT NOT NULL,
            action TEXT NOT NULL,
            entity_type TEXT NOT NULL,
            entity_id TEXT,
            details_json TEXT NOT NULL DEFAULT '{}'
        );
        CREATE INDEX IF NOT EXISTS ix_audit_created_at ON audit_log(created_at DESC);
        """)


init_db()


class KeyPair(BaseModel):
    left: str = Field(min_length=1, max_length=200)
    right: str = Field(min_length=1, max_length=200)


class ComparePair(BaseModel):
    left: str = Field(min_length=1, max_length=200)
    right: str = Field(min_length=1, max_length=200)


class ReconciliationRequest(BaseModel):
    source_a_upload_id: str
    source_b_upload_id: str
    name: str | None = Field(default=None, max_length=160)
    key_pairs: list[KeyPair] = Field(min_length=1, max_length=5)
    compare_pairs: list[ComparePair] = Field(default_factory=list, max_length=100)
    trim_whitespace: bool = True
    case_insensitive: bool = False
    numeric_tolerance: float = Field(default=0.0, ge=0, le=1_000_000_000)
    fuzzy_enabled: bool = False
    fuzzy_threshold: float = Field(default=0.85, ge=0.0, le=1.0)


class SheetSelection(BaseModel):
    sheet_name: str = Field(min_length=1, max_length=200)


class ProfileCreate(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    config: dict[str, Any]


def _audit(action: str, entity_type: str, entity_id: str | None, details: dict[str, Any] | None = None) -> None:
    # Audit metadata intentionally excludes cell values and credentials.
    with connect_db() as conn:
        conn.execute(
            "INSERT INTO audit_log(created_at, action, entity_type, entity_id, details_json) VALUES (?, ?, ?, ?, ?)",
            (utc_now(), action, entity_type, entity_id, json.dumps(details or {}, ensure_ascii=False)),
        )


def _safe_filename(name: str) -> str:
    name = Path(name or "upload.csv").name
    name = re.sub(r"[^A-Za-z0-9._ ()-]", "_", name).strip(" .")
    return name[:180] or "upload.csv"


def _validate_raw_headers(headers: list[Any] | tuple[Any, ...]) -> None:
    cleaned = [str(value).strip() if value is not None else "" for value in headers]
    if not cleaned or any(not value for value in cleaned):
        raise ValueError("Every column must have a non-empty header. Rename blank columns and upload again.")
    lowered = [value.casefold() for value in cleaned]
    if len(set(lowered)) != len(lowered):
        raise ValueError("Column names must be unique (ignoring letter case). Rename duplicate headers and upload again.")


def _validate_csv_header(path: Path) -> None:
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(8192)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                headers = next(csv.reader(handle, dialect), [])
            _validate_raw_headers(headers)
            return
        except UnicodeDecodeError:
            continue
    raise ValueError("The CSV header could not be read. Save the file as UTF-8 or a supported text encoding.")


def _validate_excel_header(path: Path, sheet_name: str) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("The selected worksheet is not present in this workbook.")
        sheet = workbook[sheet_name]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        _validate_raw_headers(list(header_row))
    finally:
        workbook.close()


def _load_frame(path: Path, extension: str, selected_sheet: str | None = None) -> tuple[pd.DataFrame, list[str]]:
    sheets: list[str] = []
    try:
        if extension == ".csv":
            _validate_csv_header(path)
            try:
                frame = pd.read_csv(path, dtype=str, keep_default_na=False, sep=None, engine="python", encoding="utf-8-sig")
            except (UnicodeDecodeError, pd.errors.ParserError):
                frame = pd.read_csv(path, dtype=str, keep_default_na=False, sep=None, engine="python", encoding="latin-1")
        elif extension in {".xlsx", ".xlsm"}:
            workbook = pd.ExcelFile(path, engine="openpyxl")
            sheets = list(workbook.sheet_names)
            if not sheets:
                raise ValueError("The Excel workbook does not contain a worksheet.")
            sheet = selected_sheet if selected_sheet in sheets else sheets[0]
            _validate_excel_header(path, sheet)
            frame = pd.read_excel(workbook, sheet_name=sheet, dtype=object)
        else:
            raise ValueError("Unsupported file format. Use CSV, XLSX, or XLSM.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"The file could not be read. Check that it is a valid {extension.upper()} file.") from exc

    if frame.empty and len(frame.columns) == 0:
        raise ValueError("The file has no header row or columns.")
    cleaned_columns = [str(column).strip() for column in frame.columns]
    if any(not column for column in cleaned_columns):
        raise ValueError("Every column must have a non-empty header. Rename blank columns and upload again.")
    lowered = [column.casefold() for column in cleaned_columns]
    if len(set(lowered)) != len(lowered):
        raise ValueError("Column names must be unique (ignoring letter case). Rename duplicate headers and upload again.")
    frame.columns = cleaned_columns
    if len(frame.index) > MAX_ROWS:
        raise ValueError(f"The file has {len(frame.index):,} rows. This MVP supports up to {MAX_ROWS:,} data rows per file.")
    return frame, sheets


def _infer_type(series: pd.Series) -> str:
    values = [v for v in series.tolist() if not is_missing(v)]
    if not values:
        return "empty"
    if pd.api.types.is_bool_dtype(series.dtype):
        return "boolean"
    if pd.api.types.is_numeric_dtype(series.dtype):
        return "number"
    text = [str(v).strip() for v in values[:500]]
    if text and all(re.fullmatch(r"[-+]?\d+(\.\d+)?", x.replace(",", "")) for x in text):
        return "number-like text"
    return "text / mixed"


def _metadata(frame: pd.DataFrame) -> dict[str, Any]:
    column_info = []
    missing_total = 0
    for column in frame.columns:
        count_missing = sum(1 for value in frame[column].tolist() if is_missing(value))
        missing_total += count_missing
        column_info.append({
            "name": str(column),
            "inferred_type": _infer_type(frame[column]),
            "missing_count": count_missing,
            "distinct_count": int(frame[column].nunique(dropna=True)) if len(frame.index) else 0,
        })
    preview_records = []
    for row in frame.head(8).to_dict(orient="records"):
        preview_records.append({str(k): json_safe(v) for k, v in row.items()})
    return {
        "row_count": int(len(frame.index)),
        "column_count": int(len(frame.columns)),
        "columns": column_info,
        "preview": preview_records,
        "missing_cells": missing_total,
    }


def _upload_by_id(upload_id: str) -> sqlite3.Row:
    with connect_db() as conn:
        row = conn.execute("SELECT * FROM uploads WHERE id = ?", (upload_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Uploaded dataset not found. Upload the file again.")
    if not Path(row["stored_path"]).is_file():
        raise HTTPException(status_code=410, detail="The stored upload is no longer available. Please upload it again.")
    return row


def _load_upload(upload_id: str) -> tuple[sqlite3.Row, pd.DataFrame]:
    row = _upload_by_id(upload_id)
    try:
        frame, _ = _load_frame(Path(row["stored_path"]), row["extension"], row["selected_sheet"])
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return row, frame


def _register_upload(filename: str, raw: bytes) -> dict[str, Any]:
    safe_name = _safe_filename(filename)
    extension = Path(safe_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Unsupported file type. Upload a .csv, .xlsx, or .xlsm file.")
    if len(raw) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="File exceeds the 50 MB upload limit.")
    if not raw:
        raise HTTPException(status_code=422, detail="The uploaded file is empty.")

    upload_id = uuid.uuid4().hex
    stored = UPLOAD_DIR / f"{upload_id}{extension}"
    stored.write_bytes(raw)
    try:
        frame, sheets = _load_frame(stored, extension)
        if len(frame.index) == 0:
            raise ValueError("The file contains headers but no data rows.")
        metadata = _metadata(frame)
    except Exception as exc:
        stored.unlink(missing_ok=True)
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(status_code=422, detail=str(exc) or "The uploaded file could not be validated.") from exc

    selected_sheet = sheets[0] if sheets else None
    with connect_db() as conn:
        conn.execute("""INSERT INTO uploads
            (id, filename, stored_path, extension, selected_sheet, sheet_names_json, columns_json,
             row_count, missing_cells, size_bytes, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (upload_id, safe_name, str(stored), extension, selected_sheet, json.dumps(sheets),
             json.dumps(metadata["columns"], ensure_ascii=False), metadata["row_count"],
             metadata["missing_cells"], len(raw), utc_now()))
    _audit("upload.created", "upload", upload_id, {
        "filename": safe_name, "rows": metadata["row_count"], "columns": metadata["column_count"],
        "format": extension,
    })
    return {
        "upload_id": upload_id,
        "filename": safe_name,
        "extension": extension,
        "size_bytes": len(raw),
        "selected_sheet": selected_sheet,
        "sheet_names": sheets,
        **metadata,
    }


def _run_summary(run: sqlite3.Row) -> dict[str, Any]:
    return {
        "id": run["id"],
        "created_at": run["created_at"],
        "name": run["name"],
        "source_a_name": run["source_a_name"],
        "source_b_name": run["source_b_name"],
        "summary": json.loads(run["summary_json"]),
        "warnings": json.loads(run["warnings_json"]),
        "config": json.loads(run["config_json"]),
    }


def _get_run(run_id: str) -> sqlite3.Row:
    with connect_db() as conn:
        row = conn.execute("SELECT * FROM runs WHERE id = ?", (run_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Reconciliation run not found.")
    return row


def _exception_page(run_id: str, page: int, page_size: int, exception_type: str | None, search: str | None) -> dict[str, Any]:
    clauses = ["run_id = ?"]
    params: list[Any] = [run_id]
    if exception_type and exception_type != "all":
        clauses.append("exception_type = ?")
        params.append(exception_type)
    if search:
        term = f"%{search[:100]}%"
        clauses.append("item_json LIKE ?")
        params.append(term)
    where = " AND ".join(clauses)
    with connect_db() as conn:
        total = conn.execute(f"SELECT COUNT(*) FROM run_exceptions WHERE {where}", params).fetchone()[0]
        rows = conn.execute(
            f"SELECT item_json FROM run_exceptions WHERE {where} ORDER BY ordinal LIMIT ? OFFSET ?",
            [*params, page_size, (page - 1) * page_size],
        ).fetchall()
    return {
        "items": [json.loads(row["item_json"]) for row in rows],
        "page": page,
        "page_size": page_size,
        "total": total,
        "pages": (total + page_size - 1) // page_size,
    }


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"status": "ok", "app": "Data Quality & Reconciliation Workbench", "version": "0.1.0"}


@app.post("/api/uploads")
async def upload_file(file: UploadFile = File(...)) -> dict[str, Any]:
    name = file.filename or "upload.csv"
    raw = await file.read(MAX_FILE_BYTES + 1)
    await file.close()
    return _register_upload(name, raw)


@app.delete("/api/uploads/{upload_id}")
def delete_upload(upload_id: str) -> dict[str, Any]:
    # Historical runs retain their summarized exceptions; only the stored source copy is removed.
    with connect_db() as conn:
        upload = conn.execute("SELECT filename, stored_path FROM uploads WHERE id = ?", (upload_id,)).fetchone()
        if upload is None:
            raise HTTPException(status_code=404, detail="Uploaded dataset not found.")
        stored_path = upload["stored_path"]
        if stored_path:
            try:
                Path(stored_path).unlink(missing_ok=True)
            except OSError as exc:
                raise HTTPException(status_code=500, detail="The saved upload could not be deleted. Check filesystem permissions.") from exc
        conn.execute("UPDATE uploads SET stored_path = '' WHERE id = ?", (upload_id,))
    _audit("upload.deleted", "upload", upload_id, {"filename": upload["filename"], "source_copy_removed": True})
    return {"deleted": True, "upload_id": upload_id, "filename": upload["filename"]}


@app.patch("/api/uploads/{upload_id}/sheet")
def select_sheet(upload_id: str, request: SheetSelection) -> dict[str, Any]:
    upload = _upload_by_id(upload_id)
    sheets = json.loads(upload["sheet_names_json"])
    if request.sheet_name not in sheets:
        raise HTTPException(status_code=422, detail="Select a worksheet listed for this workbook.")
    try:
        frame, _ = _load_frame(Path(upload["stored_path"]), upload["extension"], request.sheet_name)
        if len(frame.index) == 0:
            raise ValueError("The selected worksheet contains no data rows.")
        metadata = _metadata(frame)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    with connect_db() as conn:
        conn.execute("UPDATE uploads SET selected_sheet=?, columns_json=?, row_count=?, missing_cells=? WHERE id=?",
                     (request.sheet_name, json.dumps(metadata["columns"], ensure_ascii=False),
                      metadata["row_count"], metadata["missing_cells"], upload_id))
    _audit("upload.sheet_selected", "upload", upload_id, {"sheet": request.sheet_name})
    return {"upload_id": upload_id, "filename": upload["filename"], "extension": upload["extension"],
            "size_bytes": upload["size_bytes"], "selected_sheet": request.sheet_name,
            "sheet_names": sheets, **metadata}


@app.post("/api/demo")
def load_demo_data() -> dict[str, Any]:
    demo = {}
    for side in ("a", "b"):
        path = SAMPLE_DIR / f"source_{side}.csv"
        if not path.exists():
            raise HTTPException(status_code=500, detail="Synthetic demo files are missing from the installation.")
        demo[side] = _register_upload(path.name, path.read_bytes())
    return {"source_a": demo["a"], "source_b": demo["b"]}


@app.post("/api/reconciliations")
def create_reconciliation(request: ReconciliationRequest) -> dict[str, Any]:
    upload_a, frame_a = _load_upload(request.source_a_upload_id)
    upload_b, frame_b = _load_upload(request.source_b_upload_id)
    key_pairs = [pair.model_dump() for pair in request.key_pairs]
    compare_pairs = [pair.model_dump() for pair in request.compare_pairs]
    if len({pair["left"] for pair in key_pairs}) != len(key_pairs) or len({pair["right"] for pair in key_pairs}) != len(key_pairs):
        raise HTTPException(status_code=422, detail="Each source column can be used only once in a composite matching key.")
    if len({(pair["left"], pair["right"]) for pair in compare_pairs}) != len(compare_pairs):
        raise HTTPException(status_code=422, detail="Duplicate comparison field mappings are not allowed.")
    try:
        result = reconcile_dataframes(
            frame_a, frame_b, key_pairs, compare_pairs,
            trim_whitespace=request.trim_whitespace,
            case_insensitive=request.case_insensitive,
            numeric_tolerance=request.numeric_tolerance,
            fuzzy_enabled=request.fuzzy_enabled,
            fuzzy_threshold=request.fuzzy_threshold,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    run_id = uuid.uuid4().hex
    run_name = (request.name or f"{Path(upload_a['filename']).stem} vs {Path(upload_b['filename']).stem}").strip()[:160]
    config = result["config"]
    config["source_a_upload_id"] = request.source_a_upload_id
    config["source_b_upload_id"] = request.source_b_upload_id
    config["source_a_filename"] = upload_a["filename"]
    config["source_b_filename"] = upload_b["filename"]
    created_at = utc_now()
    with connect_db() as conn:
        conn.execute("""INSERT INTO runs
            (id, created_at, name, source_a_upload_id, source_b_upload_id, source_a_name,
             source_b_name, summary_json, config_json, warnings_json)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (run_id, created_at, run_name, request.source_a_upload_id, request.source_b_upload_id,
             upload_a["filename"], upload_b["filename"], json.dumps(result["summary"], ensure_ascii=False),
             json.dumps(config, ensure_ascii=False), json.dumps(result["warnings"], ensure_ascii=False)))
        conn.executemany("""INSERT INTO run_exceptions(run_id, ordinal, exception_type, item_json)
            VALUES (?, ?, ?, ?)""",
            [(run_id, i, item["exception_type"], json.dumps(item, ensure_ascii=False))
             for i, item in enumerate(result["exceptions"])])
    _audit("reconciliation.completed", "run", run_id, {
        "source_a": upload_a["filename"], "source_b": upload_b["filename"],
        "matched_records": result["summary"]["matched_records"],
        "exception_count": result["summary"]["exception_count"],
    })
    run = _get_run(run_id)
    payload = _run_summary(run)
    payload["exceptions"] = _exception_page(run_id, 1, 50, None, None)
    return payload


@app.get("/api/runs")
def list_runs(limit: int = Query(default=50, ge=1, le=200)) -> dict[str, Any]:
    with connect_db() as conn:
        rows = conn.execute("SELECT * FROM runs ORDER BY created_at DESC LIMIT ?", (limit,)).fetchall()
    return {"items": [_run_summary(row) for row in rows]}


@app.get("/api/runs/{run_id}")
def get_run(
    run_id: str,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    exception_type: str | None = Query(default=None, max_length=80),
    search: str | None = Query(default=None, max_length=100),
) -> dict[str, Any]:
    run = _get_run(run_id)
    payload = _run_summary(run)
    payload["exceptions"] = _exception_page(run_id, page, page_size, exception_type, search)
    return payload


@app.get("/api/runs/{run_id}/export")
def export_run(run_id: str, format: str = Query(default="xlsx", pattern="^(csv|xlsx)$")) -> StreamingResponse:
    run = _get_run(run_id)
    summary = json.loads(run["summary_json"])
    config = json.loads(run["config_json"])
    with connect_db() as conn:
        rows = conn.execute("SELECT item_json FROM run_exceptions WHERE run_id = ? ORDER BY ordinal", (run_id,)).fetchall()
    exceptions = [json.loads(row["item_json"]) for row in rows]
    columns = ["exception_type", "severity", "message", "key", "source_a_row", "source_b_row",
               "column", "source_a_value", "source_b_value", "confidence"]
    date_prefix = datetime.now().strftime("%Y%m%d-%H%M%S")
    if format == "csv":
        text_stream = io.StringIO(newline="")
        writer = csv.DictWriter(text_stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for item in exceptions:
            safe_item = {}
            for column in columns:
                value = item.get(column)
                if isinstance(value, str) and value[:1] in {"=", "+", "-", "@", "\t", "\r"}:
                    value = "'" + value
                safe_item[column] = value
            writer.writerow(safe_item)
        buffer = io.BytesIO(text_stream.getvalue().encode("utf-8-sig"))
        media_type = "text/csv; charset=utf-8"
        filename = f"reconciliation-{date_prefix}.csv"
    else:
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        summary_sheet.append(["Metric", "Value"])
        summary_sheet.append(["Run name", run["name"]])
        summary_sheet.append(["Created at (UTC)", run["created_at"]])
        summary_sheet.append(["Source A", run["source_a_name"]])
        summary_sheet.append(["Source B", run["source_b_name"]])
        for key, value in summary.items():
            summary_sheet.append([key.replace("_", " ").title(), json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
        summary_sheet.append(["Matching keys", json.dumps(config.get("key_pairs", []), ensure_ascii=False)])
        summary_sheet.append(["Compared fields", json.dumps(config.get("compare_pairs", []), ensure_ascii=False)])
        summary_sheet.freeze_panes = "A2"
        summary_sheet.column_dimensions["A"].width = 34
        summary_sheet.column_dimensions["B"].width = 88

        exception_sheet = workbook.create_sheet("Exceptions")
        exception_sheet.append(columns)
        for item in exceptions:
            values = [item.get(column) for column in columns]
            exception_sheet.append(values)
        exception_sheet.freeze_panes = "A2"
        exception_sheet.auto_filter.ref = exception_sheet.dimensions
        widths = [28, 12, 66, 32, 14, 14, 28, 30, 30, 14]
        for index, width in enumerate(widths, start=1):
            exception_sheet.column_dimensions[exception_sheet.cell(1, index).column_letter].width = width
        for sheet in workbook.worksheets:
            sheet.row_dimensions[1].height = 26
            for cell in sheet[1]:
                cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
                cell.fill = PatternFill(fill_type="solid", fgColor="172B4D")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.number_format = "@"
                        cell.data_type = "s"  # Force text: untrusted input must not become a spreadsheet formula.
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        buffer = output
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"reconciliation-{date_prefix}.xlsx"
    _audit("export.generated", "run", run_id, {"format": format, "exception_rows": len(exceptions)})
    return StreamingResponse(buffer, media_type=media_type,
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/profiles")
def list_profiles() -> dict[str, Any]:
    with connect_db() as conn:
        rows = conn.execute("SELECT * FROM profiles ORDER BY name COLLATE NOCASE").fetchall()
    return {"items": [{"id": row["id"], "name": row["name"], "config": json.loads(row["config_json"]),
                       "created_at": row["created_at"], "updated_at": row["updated_at"]} for row in rows]}


@app.post("/api/profiles")
def create_profile(request: ProfileCreate) -> dict[str, Any]:
    name = request.name.strip()
    if not name:
        raise HTTPException(status_code=422, detail="Enter a name for this rule profile.")
    config = request.config
    if not isinstance(config.get("key_pairs"), list) or not config["key_pairs"]:
        raise HTTPException(status_code=422, detail="A saved rule profile must contain at least one matching key.")
    profile_id = uuid.uuid4().hex
    now = utc_now()
    try:
        with connect_db() as conn:
            conn.execute("INSERT INTO profiles(id, name, config_json, created_at, updated_at) VALUES (?, ?, ?, ?, ?)",
                         (profile_id, name[:100], json.dumps(config, ensure_ascii=False), now, now))
    except sqlite3.IntegrityError as exc:
        raise HTTPException(status_code=409, detail="A profile with this name already exists.") from exc
    _audit("profile.created", "profile", profile_id, {"name": name[:100]})
    return {"id": profile_id, "name": name[:100], "config": config, "created_at": now, "updated_at": now}


@app.delete("/api/profiles/{profile_id}")
def delete_profile(profile_id: str) -> dict[str, Any]:
    with connect_db() as conn:
        row = conn.execute("SELECT name FROM profiles WHERE id = ?", (profile_id,)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Rule profile not found.")
        conn.execute("DELETE FROM profiles WHERE id = ?", (profile_id,))
    _audit("profile.deleted", "profile", profile_id, {"name": row["name"]})
    return {"deleted": True, "id": profile_id}


@app.get("/api/audit")
def get_audit_log(limit: int = Query(default=100, ge=1, le=500)) -> dict[str, Any]:
    with connect_db() as conn:
        rows = conn.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
    return {"items": [{"id": row["id"], "created_at": row["created_at"], "action": row["action"],
                       "entity_type": row["entity_type"], "entity_id": row["entity_id"],
                       "details": json.loads(row["details_json"])} for row in rows]}
