from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def _upload(name: str, content: str) -> str:
    response = client.post('/api/uploads', files={'file': (name, content.encode('utf-8'), 'text/csv')})
    assert response.status_code == 200, response.text
    return response.json()['upload_id']


def test_upload_reconcile_saved_profile_audit_and_exports():
    a_id = _upload('a.csv', 'id,amount\n001,10\n002,20\n')
    b_id = _upload('b.csv', 'id,amount\n001,12\n003,30\n')
    result = client.post('/api/reconciliations', json={
        'source_a_upload_id': a_id,
        'source_b_upload_id': b_id,
        'key_pairs': [{'left': 'id', 'right': 'id'}],
        'compare_pairs': [{'left': 'amount', 'right': 'amount'}],
    })
    assert result.status_code == 200, result.text
    body = result.json()
    assert body['summary']['matched_records'] == 1
    assert body['summary']['value_mismatches'] == 1
    assert body['summary']['records_only_in_a'] == 1
    run_id = body['id']

    page = client.get(f'/api/runs/{run_id}?page=1&page_size=5')
    assert page.status_code == 200
    assert page.json()['exceptions']['total'] == body['summary']['exception_count']

    profile = client.post('/api/profiles', json={'name': 'Monthly test profile', 'config': {
        'key_pairs': [{'left': 'id', 'right': 'id'}], 'compare_pairs': [{'left': 'amount', 'right': 'amount'}],
        'trim_whitespace': True, 'case_insensitive': False, 'numeric_tolerance': 0,
        'fuzzy_enabled': False, 'fuzzy_threshold': 0.85,
    }})
    assert profile.status_code == 200, profile.text
    assert client.get('/api/profiles').json()['items'][0]['name'] == 'Monthly test profile'

    csv_response = client.get(f'/api/runs/{run_id}/export?format=csv')
    xlsx_response = client.get(f'/api/runs/{run_id}/export?format=xlsx')
    assert csv_response.status_code == 200 and b'exception_type' in csv_response.content
    assert xlsx_response.status_code == 200 and xlsx_response.content.startswith(b'PK')
    assert any(event['action'] == 'reconciliation.completed' for event in client.get('/api/audit').json()['items'])

    removed = client.delete(f'/api/uploads/{a_id}')
    assert removed.status_code == 200
    # Previously computed run results remain available after the source copy is removed.
    assert client.get(f'/api/runs/{run_id}').status_code == 200


def test_upload_rejects_unsupported_extension():
    response = client.post('/api/uploads', files={'file': ('bad.exe', b'not a dataset', 'application/octet-stream')})
    assert response.status_code == 415


def test_upload_rejects_duplicate_headers():
    response = client.post('/api/uploads', files={'file': ('duplicates.csv', b'id,ID\n1,2\n', 'text/csv')})
    assert response.status_code == 422
    assert 'unique' in response.json()['detail'].lower()


def test_excel_worksheet_switch_is_supported():
    from io import BytesIO
    from openpyxl import Workbook

    workbook = Workbook()
    first = workbook.active
    first.title = 'Primary'
    first.append(['account_id', 'value'])
    first.append(['A-1', '10'])
    second = workbook.create_sheet('Alternate')
    second.append(['account_id', 'value'])
    second.append(['B-1', '20'])
    output = BytesIO()
    workbook.save(output)
    response = client.post('/api/uploads', files={'file': ('book.xlsx', output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 200, response.text
    first_result = response.json()
    assert first_result['selected_sheet'] == 'Primary'
    assert first_result['sheet_names'] == ['Primary', 'Alternate']
    switched = client.patch(f"/api/uploads/{first_result['upload_id']}/sheet", json={'sheet_name': 'Alternate'})
    assert switched.status_code == 200, switched.text
    assert switched.json()['selected_sheet'] == 'Alternate'
    assert switched.json()['preview'][0]['account_id'] == 'B-1'


def test_exports_do_not_turn_untrusted_text_into_excel_formulas():
    from io import BytesIO
    from openpyxl import load_workbook

    a_id = _upload('formula-a.csv', 'id,value\n1,=1+1\n')
    b_id = _upload('formula-b.csv', 'id,value\n1,plain\n')
    run = client.post('/api/reconciliations', json={
        'source_a_upload_id': a_id,
        'source_b_upload_id': b_id,
        'key_pairs': [{'left': 'id', 'right': 'id'}],
        'compare_pairs': [{'left': 'value', 'right': 'value'}],
    })
    assert run.status_code == 200, run.text
    run_id = run.json()['id']
    xlsx = client.get(f'/api/runs/{run_id}/export?format=xlsx')
    workbook = load_workbook(BytesIO(xlsx.content), data_only=False)
    assert workbook.sheetnames == ['Summary', 'Exceptions']
    formula_like_cells = [cell for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
                          if isinstance(cell.value, str) and cell.value == '=1+1']
    assert formula_like_cells and all(cell.data_type == 's' for cell in formula_like_cells)
    csv_export = client.get(f'/api/runs/{run_id}/export?format=csv')
    assert b"'=1+1" in csv_export.content
